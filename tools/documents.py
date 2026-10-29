# RÔLE DE CE FICHIER :
# Extrait et identifie le texte des formats documentaires autorisés. ToolRegistry
# ne rapporte une lecture réussie qu'après extraction effective du contenu.

from __future__ import annotations

import csv
import html
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from tools.filesystem import WorkspaceFilesystem

TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".yaml", ".yml", ".csv", ".html", ".htm",
    ".css", ".js", ".jsx", ".ts", ".tsx", ".py", ".xml", ".toml",
}
MAX_EXTRACTED_CHARS = 120_000
MAX_ARCHIVE_MEMBERS = 5_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 300


def _validate_archive(archive: zipfile.ZipFile) -> None:
    """Refuse les archives bureautiques dont l'extraction consommerait trop."""
    members = archive.infolist()
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise ValueError("Archive bureautique refusée : trop d'entrées.")
    if sum(member.file_size for member in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise ValueError("Archive bureautique refusée : taille décompressée trop grande.")


class _VisibleText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript"}:
            self._skip_depth += 1
        elif tag.lower() in {"p", "div", "br", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _extract_docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        _validate_archive(archive)
        xml_data = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml_data)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        text = "".join(node.text or "" for node in paragraph.iter(f"{namespace}t"))
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def _extract_xlsx(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("La lecture XLSX nécessite openpyxl.") from None
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        _validate_archive(archive)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = {}
    lines = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 1000:
                    rows.append(["[lignes suivantes tronquées]"])
                    break
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(values)
                    lines.append(f"[{sheet.title}] " + " | ".join(values))
            sheets[sheet.title] = rows
    finally:
        workbook.close()
    return "\n".join(lines), {"sheets": list(sheets), "rows": sheets}


def _extract_pptx(data: bytes) -> tuple[str, dict[str, Any]]:
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    text_ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    slides: list[tuple[int, str]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        _validate_archive(archive)
        for name in archive.namelist():
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
            if not match:
                continue
            root = ElementTree.fromstring(archive.read(name))
            chunks = [node.text or "" for node in root.iter(f"{text_ns}t")]
            slides.append((int(match.group(1)), " ".join(chunks).strip()))
    slides.sort()
    text = "\n".join(f"[Slide {index}] {content}" for index, content in slides if content)
    return text, {"slide_count": len(slides)}


def _extract_pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    try:
        import fitz
    except ImportError:
        fitz = None
    if fitz is not None:
        document = fitz.open(stream=data, filetype="pdf")
        try:
            if document.page_count > MAX_PDF_PAGES:
                raise ValueError(
                    f"PDF trop long ({document.page_count} pages; limite {MAX_PDF_PAGES})."
                )
            page_text = [page.get_text("text") for page in document]
            text = "\n".join(
                f"[Page {index + 1}] {value}"
                for index, value in enumerate(page_text)
                if value.strip()
            ).strip()
            page_count = document.page_count
            if text:
                return text, {"pages": page_count, "extraction": "embedded_text"}

            # OCR is attempted only for an image-only/scanned PDF and only when
            # both Tesseract and its Python adapter are installed.
            try:
                import pytesseract
                from PIL import Image

                pytesseract.get_tesseract_version()
            except Exception:
                raise RuntimeError(
                    "Le PDF ne contient pas de texte extractible et aucun moteur OCR "
                    "Tesseract fonctionnel n'est disponible. Le contenu n'a pas été lu."
                ) from None

            ocr_pages = min(page_count, 10)
            extracted = []
            for index in range(ocr_pages):
                pixmap = document[index].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                extracted.append(
                    f"[Page {index + 1}, OCR] "
                    + pytesseract.image_to_string(image, timeout=30)
                )
            marker = ""
            if page_count > ocr_pages:
                marker = f"\n[OCR limité aux {ocr_pages} premières pages sur {page_count}]"
            text = "\n".join(extracted).strip()
            if not text:
                raise RuntimeError("OCR exécuté, mais aucun texte lisible n'a été extrait.")
            return text + marker, {
                "pages": page_count,
                "ocr_pages": ocr_pages,
                "extraction": "ocr",
            }
        finally:
            document.close()

    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("La lecture PDF nécessite PyMuPDF ou pypdf.") from None
    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(
            f"PDF trop long ({len(reader.pages)} pages; limite {MAX_PDF_PAGES})."
        )
    page_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(
        f"[Page {index + 1}] {value}"
        for index, value in enumerate(page_text)
        if value.strip()
    ).strip()
    if not text:
        raise RuntimeError(
            "Le PDF ne contient pas de texte extractible et aucun moteur OCR n'est disponible."
        )
    return text, {"pages": len(reader.pages), "extraction": "embedded_text"}


def read_document(
    relative_path: str,
    *,
    filesystem: WorkspaceFilesystem | None = None,
    max_chars: int = MAX_EXTRACTED_CHARS,
) -> dict[str, Any]:
    """Extrait réellement un document et retourne le format et les métadonnées."""
    fs = filesystem or WorkspaceFilesystem()
    path = fs.resolve(relative_path)
    data = fs.read_bytes(relative_path)
    extension = path.suffix.lower()
    metadata: dict[str, Any] = {}

    if extension in TEXT_EXTENSIONS:
        text = _decode_text(data)
        if extension in {".json"}:
            # Vérifie que le contenu annoncé comme JSON est valide.
            json.loads(text)
        elif extension == ".csv":
            rows = list(csv.reader(io.StringIO(text)))
            text = "\n".join(" | ".join(row) for row in rows)
            metadata["rows"] = len(rows)
        elif extension in {".html", ".htm"}:
            parser = _VisibleText()
            parser.feed(text)
            text = html.unescape("".join(parser.parts))
        metadata["extraction"] = "text"
    elif extension == ".pdf":
        text, metadata = _extract_pdf(data)
    elif extension == ".docx":
        text = _extract_docx(data)
        metadata["extraction"] = "docx_xml"
    elif extension == ".xlsx":
        text, metadata = _extract_xlsx(data)
        metadata["extraction"] = "xlsx"
    elif extension == ".pptx":
        text, metadata = _extract_pptx(data)
        metadata["extraction"] = "pptx_xml"
    else:
        raise ValueError(f"Format non pris en charge : {extension or '(sans extension)'}.")

    if not text.strip():
        raise RuntimeError(f"Aucun texte n'a pu être extrait de {relative_path}.")
    text = fs.redact_secrets(text)
    original_length = len(text)
    truncated = original_length > max_chars
    if truncated:
        text = text[:max_chars] + "\n[extrait tronqué; contenu complet non transmis]"
    return {
        "path": path.relative_to(fs.root).as_posix(),
        "format": extension.lstrip("."),
        "text": text,
        "characters_extracted": original_length,
        "truncated": truncated,
        "metadata": metadata,
    }


def _verify_compiled_pdf_bytes(data: bytes) -> int:
    """Vérifie qu'un PDF compilé est valide, non vide et retourne son nombre réel de pages."""
    if not data or len(data) < 64 or not data.startswith(b"%PDF-"):
        raise ValueError("Le fichier généré n'est pas un document PDF valide (%PDF- manquant).")
    try:
        import fitz

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            pages = int(doc.page_count)
        finally:
            doc.close()
        if pages < 1:
            raise ValueError("Le PDF compilé ne contient aucune page.")
        return pages
    except ImportError:
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        if pages < 1:
            raise ValueError("Le PDF compilé ne contient aucune page.")
        return pages
    except ImportError:
        return max(1, data.count(b"/Type /Page"))


def _repair_latex_source(source: str, compiler_log: str = "") -> tuple[str, list[str]]:
    """Tente une réparation déterministe d'une source LaTeX en échec de compilation."""
    repaired = source
    notes: list[str] = []

    if "\\documentclass" not in repaired:
        repaired = (
            "\\documentclass[11pt,a4paper]{article}\n"
            "\\usepackage[utf8]{inputenc}\n"
            "\\usepackage[T1]{fontenc}\n"
            "\\usepackage{amsmath,amssymb}\n"
            "\\begin{document}\n"
            + repaired.strip()
            + "\n\\end{document}\n"
        )
        notes.append("Ajout de l'en-tête \\documentclass{article} et de l'environnement document.")
        return repaired, notes

    if "\\usepackage{libertinus-otf}" in repaired:
        repaired = repaired.replace("\\usepackage{libertinus-otf}", "% libertinus-otf désactivé")
        notes.append("Désactivation de libertinus-otf absent du cache local.")

    # Résolution des conflits unicode-math et amssymb
    if "already defined" in compiler_log and "\\usepackage{unicode-math}" in repaired and ("\\usepackage{amsmath,amssymb}" in repaired or "\\usepackage{amssymb}" in repaired):
        if "\\usepackage{amsmath,amssymb}" in repaired:
            repaired = repaired.replace("\\usepackage{amsmath,amssymb}", "")
            repaired = repaired.replace("\\usepackage{unicode-math}", "\\usepackage{amsmath,amssymb}\n\\usepackage{unicode-math}")
            notes.append("Réordonnancement de amsmath,amssymb avant unicode-math.")
        elif "\\usepackage{amssymb}" in repaired:
            repaired = repaired.replace("\\usepackage{amssymb}", "")
            repaired = repaired.replace("\\usepackage{unicode-math}", "\\usepackage{amssymb}\n\\usepackage{unicode-math}")
            notes.append("Réordonnancement de amssymb avant unicode-math.")

    # Résolution des noms de polices Latin Modern pour fontspec/Tectonic
    if "Latin Modern Roman" in repaired and ("cannot be found" in compiler_log or "fontspec" in compiler_log):
        repaired = repaired.replace("{Latin Modern Roman}", "{lmroman10-regular.otf}")
        notes.append("Adaptation du nom de police Latin Modern Roman vers lmroman10-regular.otf.")
    if "Latin Modern Math" in repaired and ("cannot be found" in compiler_log or "fontspec" in compiler_log):
        repaired = repaired.replace("{Latin Modern Math}", "{latinmodern-math.otf}")
        notes.append("Adaptation du nom de police Latin Modern Math vers latinmodern-math.otf.")

    if "\\begin{document}" not in repaired:
        lines = repaired.splitlines()
        insert_idx = 0
        for idx, line in enumerate(lines):
            if line.strip().startswith(("\\documentclass", "\\usepackage", "\\title", "\\author", "\\date", "\\newcommand", "\\setlength")):
                insert_idx = idx + 1
        lines.insert(insert_idx, "\\begin{document}")
        repaired = "\n".join(lines)
        notes.append("Ajout de \\begin{document} manquant.")

    if "\\end{document}" not in repaired:
        repaired = repaired.rstrip() + "\n\\end{document}\n"
        notes.append("Ajout de \\end{document} manquant.")

    # Réparation des signes $ non équilibrés ligne par ligne dans le corps du document
    doc_start = repaired.find("\\begin{document}")
    doc_end = repaired.rfind("\\end{document}")
    if doc_start != -1 and doc_end != -1 and doc_end > doc_start:
        preamble = repaired[: doc_start + len("\\begin{document}")]
        body = repaired[doc_start + len("\\begin{document}") : doc_end]
        tail = repaired[doc_end:]
        fixed_lines: list[str] = []
        for line in body.splitlines():
            stripped = line.lstrip()
            if stripped.startswith("%"):
                fixed_lines.append(line)
                continue
            # Compte les $ non échappés sur la ligne (hors $$)
            temp_line = re.sub(r"\\\$", "__ESC_DOLLAR__", line)
            temp_line_no_display = temp_line.replace("$$", "")
            dollar_count = temp_line_no_display.count("$")
            if dollar_count % 2 == 1:
                if any(cmd in line for cmd in ("\\int", "\\sum", "\\frac", "\\sqrt", "\\alpha", "\\beta", "\\in", "\\mathbb", "^", "_")):
                    # Insère un $ fermant ou ouvrant manquant autour de l'expression mathématique
                    math_match = re.search(r"(\\(?:int|sum|frac|sqrt|alpha|beta|mathbb)|\b[a-zA-Z0-9]+\s*[\^_])", line)
                    first_dollar = line.find("$")
                    if math_match and (first_dollar == -1 or math_match.start() < first_dollar):
                        pos = math_match.start()
                        line = line[:pos] + "$" + line[pos:]
                        notes.append("Insertion d'un délimiteur mathématique '$' ouvrant manquant.")
                    else:
                        line = line + "$"
                        notes.append("Insertion d'un délimiteur mathématique '$' fermant manquant.")
                else:
                    line = re.sub(r"(?<!\\)\$", r"\\$", line)
                    notes.append("Échappement d'un caractère '$' isolé hors mode mathématique.")
            fixed_lines.append(line)
        body = "\n".join(fixed_lines)

        # Fermeture des accolades non équilibrées dans le corps
        clean_body = re.sub(r"\\[{}]", "", body)
        open_braces = clean_body.count("{")
        close_braces = clean_body.count("}")
        if open_braces > close_braces:
            diff = open_braces - close_braces
            body = body.rstrip() + ("}" * diff) + "\n"
            notes.append(f"Fermeture automatique de {diff} accolade(s) '{{' non fermée(s).")

        # Nettoyage et échappement des '#' et titres Markdown résiduels dans le corps
        if "macro parameter character #" in compiler_log or re.search(r"(?<!\\)#", body):
            body = re.sub(r"^#{1,6}\s*(.*)$", r"\\vspace{0.3cm}\\noindent\\textbf{\\large \1}\\par\\vspace{0.15cm}", body, flags=re.MULTILINE)
            body = re.sub(r"(?<!\\)#", r"\#", body)
            notes.append("Échappement des caractères '#' et conversion des titres Markdown résiduels.")

        # Nettoyage des underscores hors mode mathématique (ex: bacblanc_sujet.tex)
        if ("Missing $ inserted" in compiler_log or "math" in compiler_log.casefold()) and "_" in body:
            parts = body.split("$")
            modified = False
            for i in range(0, len(parts), 2):
                if re.search(r"(?<!\\)_", parts[i]):
                    parts[i] = re.sub(r"(?<!\\)_", r"\_", parts[i])
                    modified = True
            if modified:
                body = "$".join(parts)
                notes.append("Échappement des caractères '_' isolés hors mode mathématique.")

        repaired = preamble + body + tail

    return repaired, notes


def compile_latex(
    tex_relative_path: str,
    *,
    filesystem: WorkspaceFilesystem | None = None,
    output_relative_path: str | None = None,
    auto_repair: bool = True,
    timeout: int = 60,
) -> dict[str, Any]:
    """Compile réellement un fichier LaTeX avec Tectonic (ou pdflatex en secours),
    vérifie la validité du PDF produit par la tentative courante et gère la réparation automatique."""
    fs = filesystem or WorkspaceFilesystem()
    tex_path = fs.resolve(tex_relative_path)
    if not tex_path.exists() or not tex_path.is_file():
        raise FileNotFoundError(f"Fichier LaTeX introuvable dans le workspace : {tex_relative_path}")

    rel_source = tex_path.relative_to(fs.root).as_posix()
    target_pdf_rel = (
        output_relative_path.strip()
        if output_relative_path and output_relative_path.strip()
        else tex_path.with_suffix(".pdf").relative_to(fs.root).as_posix()
    )
    source_text = fs.read_file(rel_source)

    tectonic_bin = shutil.which("tectonic")
    pdflatex_bin = shutil.which("pdflatex")
    compiler_name = "tectonic" if tectonic_bin else ("pdflatex" if pdflatex_bin else "")
    if not compiler_name:
        return {
            "status": "error",
            "compiler": "none",
            "source_path": rel_source,
            "pdf_path": None,
            "repaired": False,
            "repair_notes": [],
            "error": "Aucun compilateur LaTeX (tectonic ou pdflatex) n'est installé dans le PATH.",
            "log": "",
        }

    def _run_compiler_attempt(tex_content: str) -> tuple[bool, bytes | None, str]:
        with tempfile.TemporaryDirectory(prefix="neoharness_latex_") as tmp_dir_str:
            tmp_dir = Path(tmp_dir_str)
            # Copie des ressources locales éventuelles du même dossier que le .tex
            for sibling in tex_path.parent.iterdir():
                if sibling.is_file() and sibling != tex_path and sibling.suffix.lower() in {".sty", ".cls", ".png", ".jpg", ".jpeg", ".pdf", ".tex"}:
                    try:
                        shutil.copy2(sibling, tmp_dir / sibling.name)
                    except OSError:
                        pass
            work_tex = tmp_dir / "document.tex"
            work_tex.write_text(tex_content, encoding="utf-8")
            if tectonic_bin:
                cmd = [tectonic_bin, "--chatter", "minimal", "--outdir", str(tmp_dir), str(work_tex)]
            else:
                cmd = [
                    pdflatex_bin,
                    "-interaction=nonstopmode",
                    "-halt-on-error",
                    f"-output-directory={tmp_dir}",
                    str(work_tex),
                ]
            try:
                proc = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=timeout,
                    cwd=str(tmp_dir),
                )
            except subprocess.TimeoutExpired:
                return False, None, f"Délai maximal de compilation dépassé ({timeout}s)."
            except OSError as error:
                return False, None, f"Erreur d'exécution du compilateur {compiler_name} : {error}"

            log_text = "\n".join(part for part in (proc.stdout, proc.stderr) if part).strip()
            produced_pdf = tmp_dir / "document.pdf"
            if proc.returncode != 0 or not produced_pdf.is_file():
                return False, None, log_text or f"Échec de compilation (code retour {proc.returncode})."
            pdf_bytes = produced_pdf.read_bytes()
            return True, pdf_bytes, log_text

    ok, pdf_bytes, compiler_log = _run_compiler_attempt(source_text)
    repaired = False
    repair_notes: list[str] = []

    if not ok and auto_repair:
        repaired_source, notes = _repair_latex_source(source_text, compiler_log)
        if repaired_source != source_text and notes:
            ok2, pdf_bytes2, compiler_log2 = _run_compiler_attempt(repaired_source)
            if ok2 and pdf_bytes2 is not None:
                ok = True
                pdf_bytes = pdf_bytes2
                compiler_log = compiler_log2
                repaired = True
                repair_notes = notes
                fs.write_file(rel_source, repaired_source, overwrite=True)

    if not ok or pdf_bytes is None:
        return {
            "status": "error",
            "compiler": compiler_name,
            "source_path": rel_source,
            "pdf_path": None,
            "repaired": False,
            "repair_notes": repair_notes,
            "error": f"La compilation LaTeX avec {compiler_name} a échoué.",
            "log": compiler_log[-3000:],
        }

    try:
        page_count = _verify_compiled_pdf_bytes(pdf_bytes)
    except Exception as error:
        return {
            "status": "error",
            "compiler": compiler_name,
            "source_path": rel_source,
            "pdf_path": None,
            "repaired": repaired,
            "repair_notes": repair_notes,
            "error": f"PDF produit invalide : {error}",
            "log": compiler_log[-3000:],
        }

    write_info = fs.write_bytes(target_pdf_rel, pdf_bytes, overwrite=True)
    written_check = fs.read_bytes(target_pdf_rel)
    if len(written_check) != len(pdf_bytes):
        return {
            "status": "error",
            "compiler": compiler_name,
            "source_path": rel_source,
            "pdf_path": None,
            "repaired": repaired,
            "repair_notes": repair_notes,
            "error": "La vérification d'écriture du fichier PDF final a échoué.",
            "log": compiler_log[-3000:],
        }

    return {
        "status": "success",
        "compiler": compiler_name,
        "source_path": rel_source,
        "pdf_path": write_info["path"],
        "bytes": write_info["bytes"],
        "sha256": write_info["sha256"],
        "pages": page_count,
        "repaired": repaired,
        "repair_notes": repair_notes,
        "log": compiler_log[-3000:],
    }


class _HTMLStructureValidator(HTMLParser):
    VOID_TAGS = {
        "area", "base", "br", "col", "embed", "hr", "img", "input",
        "link", "meta", "param", "source", "track", "wbr", "!doctype",
    }
    CONTAINER_TAGS = {"html", "head", "body"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.stack: list[str] = []
        self.title: str = ""
        self._in_title = False
        self._current_tag = ""
        self.inline_styles: list[str] = []
        self.inline_scripts: list[str] = []
        self.linked_css: list[str] = []
        self.linked_js: list[str] = []
        self._capture_buffer: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag_lower = tag.lower()
        self._current_tag = tag_lower
        attr_dict = {k.lower(): (v or "") for k, v in attrs}

        if tag_lower == "title":
            self._in_title = True
        elif tag_lower == "link" and attr_dict.get("rel") == "stylesheet":
            href = attr_dict.get("href", "").strip()
            if href and not href.startswith(("http://", "https://", "//")):
                self.linked_css.append(href)
        elif tag_lower == "script":
            src = attr_dict.get("src", "").strip()
            if src:
                if not src.startswith(("http://", "https://", "//")):
                    self.linked_js.append(src)
            else:
                self._capture_buffer = []
        elif tag_lower == "style":
            self._capture_buffer = []

        if tag_lower not in self.VOID_TAGS and tag_lower not in self.CONTAINER_TAGS:
            self.stack.append(tag_lower)

    def handle_endtag(self, tag: str) -> None:
        tag_lower = tag.lower()
        if tag_lower == "title":
            self._in_title = False
        elif tag_lower == "script" and self._current_tag == "script":
            self.inline_scripts.append("".join(self._capture_buffer))
            self._capture_buffer = []
        elif tag_lower == "style" and self._current_tag == "style":
            self.inline_styles.append("".join(self._capture_buffer))
            self._capture_buffer = []

        if tag_lower in self.stack:
            while self.stack:
                popped = self.stack.pop()
                if popped == tag_lower:
                    break
        self._current_tag = ""

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        if self._current_tag in {"script", "style"}:
            self._capture_buffer.append(data)


def _validate_javascript_syntax(js_code: str, label: str = "script") -> tuple[bool, str]:
    """Valide réellement la syntaxe JavaScript via Node.js ou vérification de syntaxe interne."""
    trimmed = js_code.strip()
    if not trimmed:
        return True, ""

    try:
        with tempfile.NamedTemporaryFile("w", suffix=".js", encoding="utf-8", delete=False) as tmp:
            tmp.write(trimmed)
            tmp_path = tmp.name

        try:
            res = subprocess.run(
                ["node", "--check", tmp_path],
                capture_output=True,
                text=True,
                timeout=5,
            )
            if res.returncode != 0:
                err_line = res.stderr.strip()
                err_clean = err_line.replace(tmp_path, label)
                return False, f"Erreur de syntaxe JavaScript ({label}) :\n{err_clean}"
            return True, ""
        finally:
            try:
                Path(tmp_path).unlink(missing_ok=True)
            except Exception:
                pass
    except Exception:
        pass

    clean = re.sub(r"/\*[\s\S]*?\*/|//[^\n]*", "", trimmed)
    stack = []
    pairs = {")": "(", "}": "{", "]": "["}
    for char in clean:
        if char in "({[":
            stack.append(char)
        elif char in ")}]":
            if not stack or stack[-1] != pairs[char]:
                return False, f"Symbole '{char}' inattendu ou délimiteur non fermé dans {label}."
            stack.pop()
    if stack:
        return False, f"Délimiteur '{stack[-1]}' non refermé dans {label}."
    return True, ""


def _validate_css_syntax(css_code: str, label: str = "style") -> tuple[bool, str]:
    """Vérifie l'équilibre des blocs CSS et l'absence d'erreurs structurelles grossières."""
    trimmed = css_code.strip()
    if not trimmed:
        return True, ""
    clean = re.sub(r"/\*[\s\S]*?\*/", "", trimmed)
    open_braces = clean.count("{")
    close_braces = clean.count("}")
    if open_braces != close_braces:
        return False, f"Déséquilibre d'accolades dans le CSS ({label}) : {open_braces} '{{' contre {close_braces} '}}'."
    return True, ""


def _repair_html_source(source: str, unclosed_tags: list[str]) -> tuple[str, list[str]]:
    """Répare et normalise automatiquement un document HTML5 avant publication."""
    repaired = source
    notes: list[str] = []

    if not re.search(r"<!DOCTYPE\s+html", repaired, re.IGNORECASE):
        repaired = "<!DOCTYPE html>\n" + repaired.lstrip()
        notes.append("Ajout du DOCTYPE html5 manquant.")

    if not re.search(r"<html[\s>]", repaired, re.IGNORECASE):
        repaired = (
            '<!DOCTYPE html>\n<html lang="fr">\n<head>\n'
            '  <meta charset="UTF-8">\n'
            '  <meta name="viewport" content="width=device-width, initial-scale=1.0">\n'
            '  <title>Aperçu NeoHarness</title>\n'
            '</head>\n<body>\n'
            f'{repaired.strip()}\n'
            '</body>\n</html>\n'
        )
        notes.append("Encapsulation dans la structure HTML5 minimale standard (<html>, <head>, <body>).")
        return repaired, notes

    if not re.search(r"<meta[^>]+charset", repaired, re.IGNORECASE):
        head_match = re.search(r"<head[^>]*>", repaired, re.IGNORECASE)
        if head_match:
            pos = head_match.end()
            repaired = (
                repaired[:pos]
                + '\n  <meta charset="UTF-8">\n  <meta name="viewport" content="width=device-width, initial-scale=1.0">'
                + repaired[pos:]
            )
            notes.append("Insertion de la balise <meta charset=\"UTF-8\"> et viewport dans <head>.")

    if unclosed_tags:
        closing = "".join(f"</{tag}>\n" for tag in reversed(unclosed_tags))
        body_close = re.search(r"</body\s*>", repaired, re.IGNORECASE)
        if body_close:
            pos = body_close.start()
            repaired = repaired[:pos] + "\n" + closing + repaired[pos:]
        else:
            repaired = repaired.rstrip() + "\n" + closing + "</body>\n</html>\n"
        notes.append(f"Fermeture automatique des balises : {', '.join(unclosed_tags)}.")

    return repaired, notes


def compile_html(
    html_relative_path: str,
    *,
    filesystem: WorkspaceFilesystem | None = None,
    auto_repair: bool = True,
) -> dict[str, Any]:
    """Compile, valide réellement et normalise un livrable HTML/CSS/JS dans le workspace."""
    fs = filesystem or WorkspaceFilesystem()
    html_path = fs.resolve(html_relative_path)
    if not html_path.exists() or not html_path.is_file():
        raise FileNotFoundError(f"Fichier HTML introuvable dans le workspace : {html_relative_path}")

    rel_source = html_path.relative_to(fs.root).as_posix()
    source_text = fs.read_file(rel_source)
    repaired = False
    repair_notes: list[str] = []

    validator = _HTMLStructureValidator()
    try:
        validator.feed(source_text)
    except Exception as parse_err:
        if auto_repair:
            source_text, notes = _repair_html_source(source_text, [])
            repaired = True
            repair_notes.extend(notes)
            validator = _HTMLStructureValidator()
            try:
                validator.feed(source_text)
            except Exception as e2:
                return {
                    "status": "error",
                    "compiler": "html_validator",
                    "source_path": rel_source,
                    "html_path": rel_source,
                    "repaired": repaired,
                    "repair_notes": repair_notes,
                    "error": f"Erreur fatale de parsing HTML : {e2}",
                    "log": str(e2),
                }
        else:
            return {
                "status": "error",
                "compiler": "html_validator",
                "source_path": rel_source,
                "html_path": rel_source,
                "repaired": False,
                "repair_notes": [],
                "error": f"Erreur de parsing HTML : {parse_err}",
                "log": str(parse_err),
            }

    if auto_repair and (validator.stack or not re.search(r"<!DOCTYPE\s+html", source_text, re.IGNORECASE)):
        repaired_text, notes = _repair_html_source(source_text, validator.stack)
        if repaired_text != source_text:
            source_text = repaired_text
            repaired = True
            repair_notes.extend(notes)
            fs.write_file(rel_source, source_text, overwrite=True)
            validator = _HTMLStructureValidator()
            validator.feed(source_text)

    for i, css in enumerate(validator.inline_styles):
        ok, err = _validate_css_syntax(css, label=f"style inline #{i+1}")
        if not ok:
            return {
                "status": "error",
                "compiler": "html_validator",
                "source_path": rel_source,
                "html_path": rel_source,
                "repaired": repaired,
                "repair_notes": repair_notes,
                "error": err,
                "log": err,
            }

    for css_rel in validator.linked_css:
        try:
            if fs.resolve(css_rel).is_file():
                content = fs.read_file(css_rel)
                ok, err = _validate_css_syntax(content, label=f"fichier {css_rel}")
                if not ok:
                    return {
                        "status": "error",
                        "compiler": "html_validator",
                        "source_path": rel_source,
                        "html_path": rel_source,
                        "repaired": repaired,
                        "repair_notes": repair_notes,
                        "error": err,
                        "log": err,
                    }
        except Exception:
            pass

    for i, js in enumerate(validator.inline_scripts):
        ok, err = _validate_javascript_syntax(js, label=f"script inline #{i+1}")
        if not ok:
            return {
                "status": "error",
                "compiler": "html_validator",
                "source_path": rel_source,
                "html_path": rel_source,
                "repaired": repaired,
                "repair_notes": repair_notes,
                "error": err,
                "log": err,
            }

    for js_rel in validator.linked_js:
        try:
            if fs.resolve(js_rel).is_file():
                content = fs.read_file(js_rel)
                ok, err = _validate_javascript_syntax(content, label=f"fichier {js_rel}")
                if not ok:
                    return {
                        "status": "error",
                        "compiler": "html_validator",
                        "source_path": rel_source,
                        "html_path": rel_source,
                        "repaired": repaired,
                        "repair_notes": repair_notes,
                        "error": err,
                        "log": err,
                    }
        except Exception:
            pass

    raw_bytes = fs.read_bytes(rel_source)
    title = validator.title.strip() or html_path.name

    return {
        "status": "success",
        "type": "html",
        "compiler": "html_validator",
        "source_path": rel_source,
        "html_path": rel_source,
        "title": title,
        "bytes": len(raw_bytes),
        "repaired": repaired,
        "repair_notes": repair_notes,
        "has_css": bool(validator.inline_styles or validator.linked_css),
        "has_js": bool(validator.inline_scripts or validator.linked_js),
        "scripts_checked": len(validator.inline_scripts) + len(validator.linked_js),
        "styles_checked": len(validator.inline_styles) + len(validator.linked_css),
        "dependencies": validator.linked_css + validator.linked_js,
        "log": f"Compilation et validation HTML5 réussie ({len(raw_bytes)} octets, JS validé, CSS validé).",
    }


# ---------------------------------------------------------------------------
# CHANTIER 7 : DATA ANALYST AUTONOME (PROFILAGE STATISTIQUE & EXCEL .XLSX ENRICHI)
# ---------------------------------------------------------------------------

def _try_float(val: Any) -> float | None:
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("\u00a0", "").replace(" ", "")
    if not s or s.startswith("="):
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def analyze_dataset(
    relative_path: str,
    *,
    sheet_name: str = "",
    filesystem: WorkspaceFilesystem | None = None,
) -> dict[str, Any]:
    """Analyse statistique autonome d'un jeu de données (.csv, .tsv, .xlsx, .json, .sqlite/.db) :
    schéma, valeurs manquantes, moyenne/médiane/écart-type, corrélations de Pearson et détection d'anomalies (outliers)."""
    import hashlib
    import math
    import sqlite3
    import statistics

    fs = filesystem or WorkspaceFilesystem()
    target = fs.resolve(relative_path)
    if not target.is_file():
        raise FileNotFoundError(f"Jeu de données introuvable : {relative_path}")

    rel = target.relative_to(fs.root).as_posix()
    ext = target.suffix.lower()
    headers: list[str] = []
    rows: list[list[Any]] = []

    if ext in {".csv", ".tsv"}:
        raw_text = _decode_text(fs.read_bytes(rel))
        delim = "\t" if ext == ".tsv" else (";" if raw_text.splitlines()[0].count(";") > raw_text.splitlines()[0].count(",") else ",")
        all_rows = [r for r in csv.reader(io.StringIO(raw_text), delimiter=delim) if any( str(c).strip() for c in r )]
        if all_rows:
            headers = [str(h).strip() or f"col_{i+1}" for i, h in enumerate(all_rows[0])]
            rows = all_rows[1:]
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(fs.read_bytes(rel)), read_only=True, data_only=True)
        try:
            ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.worksheets[0]
            all_rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c is not None and str(c).strip() != "" for c in r)]
            if all_rows:
                headers = [str(h).strip() if h is not None else f"col_{i+1}" for i, h in enumerate(all_rows[0])]
                rows = all_rows[1:]
        finally:
            wb.close()
    elif ext == ".json":
        parsed = json.loads(_decode_text(fs.read_bytes(rel)))
        records = parsed if isinstance(parsed, list) else (parsed.get("data") or parsed.get("items") or [])
        if isinstance(records, list) and records and isinstance(records[0], dict):
            headers = list(records[0].keys())
            rows = [[rec.get(h) for h in headers] for rec in records if isinstance(rec, dict)]
    elif ext in {".sqlite", ".db"}:
        con = sqlite3.connect(str(target))
        try:
            tables = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
            t_name = sheet_name if (sheet_name and sheet_name in tables) else (tables[0] if tables else "")
            if t_name:
                cur = con.execute(f'SELECT * FROM "{t_name}" LIMIT 5000')
                headers = [d[0] for d in cur.description]
                rows = [list(r) for r in cur.fetchall()]
        finally:
            con.close()
    else:
        raise ValueError(f"Format de données non supporté pour analyze_dataset : {ext}")

    row_count = len(rows)
    col_count = len(headers)
    columns_profile: list[dict[str, Any]] = []
    numeric_series: dict[str, dict[int, float]] = {}

    for col_idx, col_name in enumerate(headers):
        raw_vals = [r[col_idx] if col_idx < len(r) else None for r in rows]
        non_empty = [v for v in raw_vals if v is not None and str(v).strip() != ""]
        missing = row_count - len(non_empty)
        num_map: dict[int, float] = {}
        for r_idx, v in enumerate(raw_vals):
            fv = _try_float(v)
            if fv is not None and math.isfinite(fv):
                num_map[r_idx] = fv

        if len(num_map) >= max(1, int(0.6 * len(non_empty))) and len(num_map) > 0:
            vals = list(num_map.values())
            vals_sorted = sorted(vals)
            n_v = len(vals)
            mean_v = statistics.mean(vals)
            median_v = statistics.median(vals_sorted)
            stdev_v = statistics.stdev(vals) if n_v > 1 else 0.0
            q1 = vals_sorted[n_v // 4]
            q3 = vals_sorted[(3 * n_v) // 4]
            iqr = q3 - q1

            outliers: list[dict[str, Any]] = []
            for r_idx, fv in num_map.items():
                is_z_outlier = stdev_v > 0 and abs(fv - mean_v) / stdev_v > 2.4
                is_iqr_outlier = iqr > 0 and (fv < q1 - 1.5 * iqr or fv > q3 + 1.5 * iqr)
                if is_z_outlier or is_iqr_outlier:
                    outliers.append({"row": r_idx + 2, "value": round(fv, 4)})

            numeric_series[col_name] = num_map
            columns_profile.append({
                "name": col_name,
                "type": "numeric",
                "count": n_v,
                "missing": missing,
                "min": round(min(vals), 4),
                "max": round(max(vals), 4),
                "sum": round(sum(vals), 4),
                "mean": round(mean_v, 4),
                "median": round(median_v, 4),
                "std_dev": round(stdev_v, 4),
                "outliers_count": len(outliers),
                "outliers": outliers[:8],
            })
        else:
            freq: dict[str, int] = {}
            for v in non_empty:
                sv = str(v).strip()[:80]
                freq[sv] = freq.get(sv, 0) + 1
            top_vals = sorted(freq.items(), key=lambda x: x[1], reverse=True)[:5]
            columns_profile.append({
                "name": col_name,
                "type": "categorical",
                "count": len(non_empty),
                "missing": missing,
                "unique_values": len(freq),
                "top_values": [{"value": k, "count": c} for k, c in top_vals],
            })

    # Calcul des corrélations de Pearson entre colonnes numériques
    correlations: list[dict[str, Any]] = []
    num_cols = list(numeric_series.keys())
    for i in range(len(num_cols)):
        for j in range(i + 1, len(num_cols)):
            c1, c2 = num_cols[i], num_cols[j]
            common_idx = sorted(set(numeric_series[c1].keys()) & set(numeric_series[c2].keys()))
            if len(common_idx) >= 3:
                x = [numeric_series[c1][k] for k in common_idx]
                y = [numeric_series[c2][k] for k in common_idx]
                mx, my = statistics.mean(x), statistics.mean(y)
                num = sum((a - mx) * (b - my) for a, b in zip(x, y))
                den = math.sqrt(sum((a - mx) ** 2 for a in x) * sum((b - my) ** 2 for b in y))
                if den > 1e-12:
                    r_val = round(num / den, 4)
                    correlations.append({"col_a": c1, "col_b": c2, "pearson_r": r_val})

    correlations.sort(key=lambda item: abs(item["pearson_r"]), reverse=True)
    return {
        "status": "success",
        "verified": True,
        "path": rel,
        "format": ext.lstrip("."),
        "row_count": row_count,
        "column_count": col_count,
        "columns": columns_profile,
        "correlations": correlations[:15],
        "sample_rows": rows[:5],
    }


def generate_excel_workbook(
    output_path: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    sheet_title: str = "Données",
    include_formulas: bool = True,
    filesystem: WorkspaceFilesystem | None = None,
) -> dict[str, Any]:
    """Génère un classeur Excel (.xlsx) professionnel stylisé avec calcul automatique des totaux/moyennes par formules natives."""
    import hashlib
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    fs = filesystem or WorkspaceFilesystem()
    clean_out = (output_path or "analyse_donnees.xlsx").strip()
    if not clean_out.lower().endswith(".xlsx"):
        clean_out += ".xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Données")[:31]

    header_fill = PatternFill(start_color="185ABD", end_color="185ABD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    total_fill = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")
    total_font = Font(color="0F6CBD", bold=True, size=10.5)
    thin_border = Border(
        left=Side(style="thin", color="D2D2D2"),
        right=Side(style="thin", color="D2D2D2"),
        top=Side(style="thin", color="D2D2D2"),
        bottom=Side(style="thin", color="D2D2D2"),
    )

    ws.append(list(headers))
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    numeric_cols: set[int] = set()
    for r_data in rows:
        converted_row: list[Any] = []
        for c_idx, val in enumerate(r_data):
            fv = _try_float(val)
            if fv is not None and not isinstance(val, str):
                converted_row.append(fv)
                numeric_cols.add(c_idx + 1)
            elif fv is not None and isinstance(val, str) and re.fullmatch(r"-?\d+(?:[.,]\d+)?", val.strip()):
                converted_row.append(fv)
                numeric_cols.add(c_idx + 1)
            else:
                converted_row.append(val)
        ws.append(converted_row)

    last_data_row = len(rows) + 1
    formulas_added = 0
    if include_formulas and len(rows) >= 2 and numeric_cols:
        sum_row_idx = last_data_row + 1
        avg_row_idx = last_data_row + 2
        ws.cell(row=sum_row_idx, column=1, value="TOTAL").font = total_font
        ws.cell(row=avg_row_idx, column=1, value="MOYENNE").font = total_font
        for c_idx in range(1, len(headers) + 1):
            col_let = get_column_letter(c_idx)
            c_sum = ws.cell(row=sum_row_idx, column=c_idx)
            c_avg = ws.cell(row=avg_row_idx, column=c_idx)
            c_sum.fill = total_fill
            c_avg.fill = total_fill
            c_sum.border = thin_border
            c_avg.border = thin_border
            if c_idx in numeric_cols:
                c_sum.value = f"=SUM({col_let}2:{col_let}{last_data_row})"
                c_avg.value = f"=AVERAGE({col_let}2:{col_let}{last_data_row})"
                c_sum.font = total_font
                c_avg.font = total_font
                formulas_added += 2

    # Ajustement automatique de la largeur des colonnes
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 42))

    buf = io.BytesIO()
    wb.save(buf)
    raw_bytes = buf.getvalue()
    write_info = fs.write_bytes(clean_out, raw_bytes, overwrite=True)

    return {
        "status": "success",
        "verified": True,
        "path": write_info["path"],
        "size_bytes": write_info["bytes"],
        "sha256": hashlib.sha256(raw_bytes).hexdigest(),
        "rows_written": len(rows),
        "columns_written": len(headers),
        "formulas_added": formulas_added,
    }

